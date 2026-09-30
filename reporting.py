"""Generate an editable, local DOCX diagnosis report."""

from __future__ import annotations

import io
from datetime import date

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "Arial Unicode MS", "DejaVu Sans"]
plt.rcParams["axes.unicode_minus"] = False


def _font(run, size: int = 10, bold: bool = False) -> None:
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def _paragraph(doc: Document, text: str, bold_prefix: str = "") -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    if bold_prefix:
        _font(paragraph.add_run(bold_prefix), 10, True)
    _font(paragraph.add_run(text), 10)


def _table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for index, value in enumerate(headers):
        table.rows[0].cells[index].text = value
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)
    for row_index, row in enumerate(table.rows):
        row_properties = row._tr.get_or_add_trPr()
        no_split = OxmlElement("w:cantSplit")
        row_properties.append(no_split)
        if row_index == 0:
            repeated_header = OxmlElement("w:tblHeader")
            repeated_header.set(qn("w:val"), "true")
            row_properties.append(repeated_header)
        for cell in row.cells:
            properties = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                line = OxmlElement(f"w:{edge}")
                line.set(qn("w:val"), "single")
                line.set(qn("w:sz"), "4")
                line.set(qn("w:color"), "D9D9D9")
                borders.append(line)
            properties.append(borders)
            if row_index == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "EAF0F5")
                properties.append(shading)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.space_before = Pt(2)
                for run in paragraph.runs:
                    _font(run, 9, row_index == 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def _chart(categories: dict) -> io.BytesIO | None:
    if len(categories) < 2:
        return None
    labels = list(categories)
    values = [round(categories[name]["ratio"] * 100) for name in labels]
    fig = plt.figure(figsize=(8.0, 3.6))
    ax = fig.add_subplot(1, 2, 1)
    bars = ax.barh(labels[::-1], values[::-1], color="#3978a8")
    ax.set_xlim(0, 110)
    ax.set_xlabel("简历证据覆盖率（%）", fontsize=10)
    ax.tick_params(labelsize=9)
    for bar, value in zip(bars, values[::-1]):
        ax.text(min(value + 2, 104), bar.get_y() + bar.get_height() / 2, f"{value}%", va="center", fontsize=9)
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="x", alpha=0.2)
    ax.set_axisbelow(True)
    if len(categories) >= 3:
        radar = fig.add_subplot(1, 2, 2, projection="polar")
        angles = np.linspace(0, 2 * np.pi, len(labels), endpoint=False).tolist()
        closed_values = values + values[:1]
        closed_angles = angles + angles[:1]
        radar.plot(closed_angles, closed_values, color="#3978a8", linewidth=2)
        radar.fill(closed_angles, closed_values, color="#3978a8", alpha=0.18)
        radar.set_xticks(angles, labels, fontsize=9)
        radar.set_ylim(0, 100)
        radar.set_yticks([25, 50, 75, 100], ["25", "50", "75", "100"], fontsize=7)
        radar.grid(alpha=0.35)
    fig.tight_layout()
    output = io.BytesIO()
    fig.savefig(output, format="png", dpi=150)
    plt.close(fig)
    output.seek(0)
    return output


def build_report(analysis: dict, student_label: str = "学生（匿名）") -> bytes:
    profile = analysis["profile"]
    target = analysis["target"]
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2.3)
    styles = doc.styles
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.color.rgb = RGBColor(0, 0, 0)
    styles["Normal"].font.size = Pt(10)
    styles["Title"].font.size = Pt(18)
    styles["Heading 1"].font.size = Pt(13)
    styles["Heading 2"].font.size = Pt(11)
    title_properties = styles["Title"]._element.get_or_add_pPr()
    old_border = title_properties.find(qn("w:pBdr"))
    if old_border is not None:
        title_properties.remove(old_border)
    no_border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "nil")
    no_border.append(bottom)
    title_properties.append(no_border)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(title.add_run("机械类学生简历诊断与岗位匹配报告"), 18, True)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(meta.add_run(f"{student_label}  |  目标岗位：{target['title']}  |  {date.today().isoformat()}"), 9)
    _paragraph(doc, analysis["disclaimer"])

    doc.add_heading("一 学生简历画像", level=1)
    _table(doc, ["维度", "从简历提取的信息"], [
        ["学历层次", profile.degree or "未识别"],
        ["专业", "、".join(profile.majors) or "未识别"],
        ["专业技能", "、".join(profile.skills) or "未识别"],
        ["工具应用", "、".join(profile.tools) or "未识别"],
        ["工作或实习", f"识别到 {len(profile.work_lines)} 行经历描述"],
        ["项目经验", f"识别到 {len(profile.project_lines)} 行经历描述"],
        ["成果证据", f"识别到 {len(profile.outcome_lines)} 行可能包含结果的描述，需人工核对"],
    ])
    for warning in profile.warnings:
        _paragraph(doc, warning, "识别提示：")

    doc.add_heading("二 目标岗位需求与匹配", level=1)
    _paragraph(doc, f"岗位来源：{target['source']}；企业：{target['employer']}。")
    requirement_rows = [[name, "、".join(value) if isinstance(value, list) else (value or "未提供")]
                        for name, value in target["requirements"].items()]
    _table(doc, ["岗位需求维度", "当前输入的岗位要求"], requirement_rows)
    if target["description"]:
        _paragraph(doc, target["description"][:500], "岗位说明摘录：")
    score_label = f"{target['score']} 分" if target["score"] is not None else "岗位要求不足，暂不评分"
    _paragraph(doc, f"简历证据匹配度：{score_label}。该分数仅对当前输入的岗位要求和简历文本有效。")
    rows = []
    for name, info in target["categories"].items():
        rows.append([name, f"{info['weight']}%", f"{info['ratio'] * 100:.0f}%",
                     "、".join(info["matched"]) or "未识别", "、".join(info["missing"]) or "无"])
    if rows:
        _table(doc, ["维度", "权重", "覆盖率", "已找到证据", "未在简历找到"], rows)
    chart = _chart(target["categories"])
    if chart:
        doc.add_picture(chart, width=Cm(15.5))

    doc.add_heading("三 分维度指标与优劣势诊断", level=1)
    _paragraph(doc, "以下‘优势’与‘待改进’均指当前简历呈现的书面证据，不等同于学生真实能力。覆盖率=已找到要求数/岗位要求数；加权得分=本岗位归一后的权重×覆盖率。")
    assessment = analysis["assessment"]
    _table(doc, ["指标", "覆盖率", "本岗位权重", "加权得分", "判断"], [
        [item["维度"], f"{item['覆盖率']}%", f"{item['权重']}%", f"{item['加权得分']} 分", item["状态"]]
        for item in assessment["指标"]
    ])
    doc.add_heading("已展示的优势", level=2)
    for line in assessment["优势"] or ["当前简历尚无达到‘文本覆盖较高’阈值的岗位维度；需结合真实经历核对。"]:
        _paragraph(doc, line)
    doc.add_heading("待改进与证据缺口", level=2)
    for line in assessment["待改进"] or ["当前规则未识别到明显缺口；仍需人工检查证据的真实性与深度。"]:
        _paragraph(doc, line)
    doc.add_heading("指标对应的简历原句", level=2)
    for item in assessment["指标"]:
        evidence = item["证据原句"]
        if evidence:
            _paragraph(doc, "；".join(f"{term}：{line}" for term, line in list(evidence.items())[:5]), f"{item['维度']}：")
        else:
            _paragraph(doc, "未在相应栏目找到直接原句。", f"{item['维度']}：")

    doc.add_heading("四 工作职责 项目经历与成果", level=1)
    for label, items in (("工作/实习职责", profile.work_lines), ("项目经历", profile.project_lines),
                         ("成果线索", profile.outcome_lines)):
        doc.add_heading(label, level=2)
        if items:
            for line in items[:8]:
                _paragraph(doc, line)
            if len(items) > 8:
                _paragraph(doc, f"其余 {len(items) - 8} 行未在报告中展开，原简历仍应人工复核。")
        else:
            _paragraph(doc, "未从相应栏目识别到文本；这表示书面证据不足，不等于没有实际经历。")

    doc.add_heading("五 原句诊断与修改建议", level=1)
    for index, item in enumerate(analysis["suggestions"], 1):
        doc.add_heading(f"建议 {index}  {item['位置']} · {item['类型']}", level=2)
        _paragraph(doc, item["原文"], "原文或现状：")
        _paragraph(doc, item["问题"], "问题：")
        if item["建议替换句"]:
            _paragraph(doc, item["建议替换句"], "建议替换句：")
        else:
            _paragraph(doc, "原句事实不足，暂不能安全生成直接替换句。", "建议替换句：")
        _paragraph(doc, item["需核实"], "核实与补充：")
    if not analysis["suggestions"]:
        _paragraph(doc, "当前规则未发现明确的结构问题；仍建议由指导教师核对事实、表达与投递岗位。")

    doc.add_heading("六 岗位推荐与下一步行动", level=1)
    ranked = analysis["ranked"][:5]
    _table(doc, ["排序", "岗位", "企业", "书面证据匹配度", "来源"], [
        [str(i), item["title"], item["employer"], f"{item['score']} 分" if item["score"] is not None else "未评分",
         "演示样例" if item["source"].startswith("演示") else item["source"]]
        for i, item in enumerate(ranked, 1)
    ])
    gaps = target["gaps"][:4]
    if gaps:
        _paragraph(doc, "优先核实以下缺口是否已有真实经历；有则补充证据，没有则安排课程、实训或项目实践：" + "；".join(gaps) + "。")
    _paragraph(doc, "报告需与真实岗位说明、简历原件及面试核实结合使用。示例岗位不能用于企业录用判断；系统尚未完成企业数据采集、专家权重校准及盲测准确率验证。")
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()
