from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfWriter

from reporting import build_report
from resume_engine import (
    analyze_resume, job_from_description, load_demo_jobs, load_jobs_csv,
    parse_resume, read_document,
)


ROOT = Path(__file__).resolve().parents[1]
SAMPLE = (ROOT / "samples" / "演示简历.txt").read_text(encoding="utf-8")


def test_resume_profile_extracts_evidence_without_inventing_outcomes():
    profile = parse_resume(SAMPLE)
    assert profile.degree == "本科"
    assert "机械设计制造及其自动化" in profile.majors
    assert "SolidWorks" in profile.tools
    assert "AutoCAD" in profile.tools
    assert profile.project_lines
    assert any("2 份" in line for line in profile.outcome_lines)
    assert not any("提升 80%" in line for line in profile.outcome_lines)


def test_matching_ranks_relevant_job_and_preserves_missing_evidence():
    jobs = load_demo_jobs(ROOT / "data" / "demo_jobs.json")
    analysis = analyze_resume(SAMPLE, jobs, selected_index=0)
    assert analysis["ranked"][0]["title"] == "机械设计工程师"
    assert analysis["target"]["score"] is not None
    assert "ANSYS" in analysis["target"]["categories"]["工具应用"]["missing"]
    assert "录用概率" in analysis["disclaimer"]
    assert any("不等于不会" in item["问题"] for item in analysis["suggestions"])


def test_docx_txt_and_scanned_pdf_handling():
    assert "机械设计" in read_document("sample.txt", SAMPLE.encode("utf-8"))
    source = Document()
    source.add_paragraph(SAMPLE)
    file = BytesIO()
    source.save(file)
    assert "SolidWorks" in read_document("sample.docx", file.getvalue())
    blank_pdf = PdfWriter()
    blank_pdf.add_blank_page(width=200, height=200)
    file = BytesIO()
    blank_pdf.write(file)
    try:
        read_document("scan.pdf", file.getvalue())
    except ValueError as exc:
        assert "OCR" in str(exc)
    else:
        raise AssertionError("An image-only PDF must not be treated as parsed text")


def test_csv_import_and_pasted_job_requirements():
    content = (ROOT / "data" / "jobs_template.csv").read_bytes()
    jobs = load_jobs_csv(content)
    assert jobs[0].title == "机械设计工程师"
    assert "SolidWorks" in jobs[0].tools
    pasted = job_from_description("自动化工程师", "本科及以上，自动化专业，熟悉 PLC 和 Python。", employer="某企业")
    assert pasted.min_degree == "本科"
    assert "PLC" in pasted.tools
    assert "Python" in pasted.tools
    sparse = job_from_description("信息不足岗位", "本科，熟悉 CAD。")
    sparse_result = analyze_resume(SAMPLE, [sparse])["target"]
    assert sparse_result["score"] is None
    assert "少于 3 个" in sparse_result["warning"]
    family_job = job_from_description("机械岗位", "本科，机械类专业；机械设计、三维建模、SolidWorks。")
    family_result = analyze_resume(SAMPLE, [family_job])["target"]
    assert family_result["categories"]["专业"]["ratio"] == 1.0


def test_report_is_editable_docx_with_requested_sections():
    jobs = load_demo_jobs(ROOT / "data" / "demo_jobs.json")
    analysis = analyze_resume(SAMPLE, jobs, selected_index=0)
    data = build_report(analysis, "虚构演示学生")
    doc = Document(BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "简历诊断与岗位匹配报告" in text
    assert "工作职责 项目经历与成果" in text
    assert "原句诊断与修改建议" in text
    assert "分维度指标与优劣势诊断" in text
    assert "建议替换句" in text
    assert "虚构演示学生" in text


def test_sentence_rewrite_preserves_documented_facts_and_assessment_explains_gaps():
    jobs = load_demo_jobs(ROOT / "data" / "demo_jobs.json")
    analysis = analyze_resume(SAMPLE, jobs, selected_index=0)
    rewrites = [item for item in analysis["suggestions"] if item["建议替换句"]]
    assert rewrites
    project = next(item for item in rewrites if "SolidWorks" in item["原文"])
    assert "SolidWorks" in project["建议替换句"]
    assert "2 份装配图" in project["建议替换句"]
    assert "80%" not in project["建议替换句"]
    assessment = analysis["assessment"]
    assert any(item["维度"] == "工具应用" and "ANSYS" in item["未找到"] for item in assessment["指标"])
    assert any("应用深度" in line for line in assessment["待改进"])
